from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path("src/data/templates/referencia-elogiosa-base.docx")
CREST = Path("public/brasao-1-bimec.png")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn("w:" + key), str(value))


document = Document()
section = document.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)

normal = document.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

header = document.add_table(rows=1, cols=2)
header.alignment = WD_TABLE_ALIGNMENT.CENTER
header.autofit = False
header.columns[0].width = Cm(2.4)
header.columns[1].width = Cm(13.5)
for cell in header.rows[0].cells:
    set_cell_border(cell, top={"val": "nil"}, left={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})
if CREST.exists():
    image_paragraph = header.cell(0, 0).paragraphs[0]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(CREST), width=Cm(1.8))
title_cell = header.cell(0, 1)
title_cell.vertical_alignment = 1
title_paragraph = title_cell.paragraphs[0]
title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_paragraph.add_run("1º BATALHÃO DE INFANTARIA MECANIZADO (ESCOLA)")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(12)
subtitle = title_cell.add_paragraph("REGIMENTO SAMPAIO")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True
subtitle.runs[0].font.name = "Arial"
subtitle.runs[0].font.size = Pt(11)

document.add_paragraph("")
heading = document.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
heading.paragraph_format.space_before = Pt(6)
heading.paragraph_format.space_after = Pt(14)
heading_run = heading.add_run("REFERÊNCIA ELOGIOSA")
heading_run.bold = True
heading_run.underline = True
heading_run.font.name = "Arial"
heading_run.font.size = Pt(14)

recipient = document.add_paragraph()
recipient.alignment = WD_ALIGN_PARAGRAPH.CENTER
recipient.paragraph_format.space_after = Pt(14)
recipient_run = recipient.add_run("{{DESTINATARIO}}")
recipient_run.bold = True
recipient_run.font.name = "Arial"
recipient_run.font.size = Pt(12)

body = document.add_paragraph()
body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
body.paragraph_format.first_line_indent = Cm(1.25)
body.paragraph_format.line_spacing = 1.15
body.paragraph_format.space_after = Pt(8)
body_run = body.add_run("{{TEXTO}}")
body_run.font.name = "Arial"
body_run.font.size = Pt(12)

motto = document.add_paragraph()
motto.alignment = WD_ALIGN_PARAGRAPH.CENTER
motto.paragraph_format.space_before = Pt(10)
motto_run = motto.add_run("Leões de Guerra! Aço! SAMPAIO!")
motto_run.bold = True
motto_run.font.name = "Arial"
motto_run.font.size = Pt(11)

date = document.add_paragraph()
date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
date.paragraph_format.space_before = Pt(12)
date_run = date.add_run("{{DATA}}")
date_run.font.name = "Arial"
date_run.font.size = Pt(12)

document.add_paragraph("")
signature = document.add_paragraph()
signature.alignment = WD_ALIGN_PARAGRAPH.CENTER
signature_run = signature.add_run("ALEX FERREIRA GOMES JÚNIOR – Ten Cel")
signature_run.bold = True
signature_run.font.name = "Arial"
signature_run.font.size = Pt(11)
function = document.add_paragraph("Comandante do 1º Batalhão de Infantaria Mecanizado (Escola)")
function.alignment = WD_ALIGN_PARAGRAPH.CENTER
function.runs[0].font.name = "Arial"
function.runs[0].font.size = Pt(11)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("1º BIMEC (Es) – Regimento Sampaio")
footer_run.font.name = "Arial"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = None

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
