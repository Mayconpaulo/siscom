from docx import Document
from pathlib import Path
from zipfile import BadZipFile
from lxml import etree

path = Path("src/data/templates/referencia-elogiosa-modelo.docx")
try:
    document = Document(path)
except BadZipFile:
    xml_path = Path("tmp-docx-extracted/word/document.xml")
    root = etree.parse(str(xml_path))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    body = root.find("w:body", namespaces)
    for index, element in enumerate(body):
        tag = etree.QName(element).localname
        if tag == "p":
            text = "".join(element.xpath(".//w:t/text()", namespaces=namespaces))
            tabs = len(element.xpath(".//w:tab", namespaces=namespaces))
            breaks = len(element.xpath(".//w:br", namespaces=namespaces))
            if text.strip() or tabs or breaks:
                print("XML-P", index, repr(text), "tabs=", tabs, "breaks=", breaks)
        elif tag == "tbl":
            rows = []
            for row in element.xpath("./w:tr", namespaces=namespaces):
                cells = []
                for cell in row.xpath("./w:tc", namespaces=namespaces):
                    cells.append("".join(cell.xpath(".//w:t/text()", namespaces=namespaces)))
                rows.append(cells)
            print("XML-TABLE", index, rows)
    raise SystemExit(0)

print("SECTIONS", len(document.sections), "PARAGRAPHS", len(document.paragraphs), "TABLES", len(document.tables))
for index, section in enumerate(document.sections):
    print("SECTION", index, "margins", section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)
    for label, part in (("HEADER", section.header), ("FOOTER", section.footer)):
        for paragraph_index, paragraph in enumerate(part.paragraphs):
            if paragraph.text.strip():
                print(label, paragraph_index, repr(paragraph.text), paragraph.style.name, paragraph.alignment)

for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.replace("\t", "<TAB>").replace("\n", "<NL>")
    if not text.strip():
        continue
    runs = [
        (run.text, run.bold, run.italic, run.underline, run.font.name, run.font.size.pt if run.font.size else None)
        for run in paragraph.runs
    ]
    print("P", index, repr(text), "style=", paragraph.style.name, "align=", paragraph.alignment, "runs=", runs)

for table_index, table in enumerate(document.tables):
    print("TABLE", table_index, len(table.rows), len(table.columns))
    for row_index, row in enumerate(table.rows):
        print("ROW", row_index, [cell.text for cell in row.cells])
